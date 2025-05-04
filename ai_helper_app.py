
import streamlit as st
import base64
import os
from PIL import Image
import pdfplumber
import docx
import openpyxl
import mimetypes

# Configurazione
st.set_page_config(layout="wide", page_title="AI Helper", page_icon="📩")

# --- Funzioni ---

def read_pdf(file):
    text = ""
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            text += page.extract_text() + "\n"
    return text

def read_docx(file):
    doc = docx.Document(file)
    return "\n".join([p.text for p in doc.paragraphs])

def read_excel(file):
    wb = openpyxl.load_workbook(file, data_only=True)
    sheet = wb.active
    text = ""
    for row in sheet.iter_rows(values_only=True):
        text += " | ".join([str(cell) if cell is not None else "" for cell in row]) + "\n"
    return text

def get_file_text(uploaded_files):
    combined_text = ""
    for uploaded_file in uploaded_files:
        file_type = uploaded_file.type
        if file_type == "application/pdf":
            combined_text += read_pdf(uploaded_file)
        elif file_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            combined_text += read_docx(uploaded_file)
        elif file_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            combined_text += read_excel(uploaded_file)
        elif file_type.startswith("text/"):
            combined_text += str(uploaded_file.read(), "utf-8")
        else:
            combined_text += f"\n[File non supportato: {uploaded_file.name}]"
    return combined_text

# --- Layout ---

# Logo
st.image("logo.png", width=180)

# Colonne layout
col1, col2 = st.columns([1, 1])

with col1:
    st.markdown("## Nuova Analisi")
    st.write("Salve! Posso aiutarti ad analizzare e strutturare una discussione email concentrata su dettagli tecnici e commerciali. Se hai un thread di email da esaminare, sentiti libero di copiarle+incollarle qui. Una volta avviata l'analisi, ti fornirò un riepilogo cronologico con codici, tecnici, azioni principali e prossimi step.")
    
    email_text = st.text_area("Inputo dati", "Incolla qui il contenuto dell’email o testo da analizzare", height=150)

    uploaded_files = st.file_uploader("Allega file (PDF, Word, Excel, TXT)", accept_multiple_files=True, type=["pdf", "docx", "xlsx", "txt"])

    if st.button("Nuova Analisi"):
        input_text = email_text
        if uploaded_files:
            input_text += "\n\n" + get_file_text(uploaded_files)
        st.session_state["input_text"] = input_text
        st.session_state["output_text"] = "🧠 Elaborazione AI in corso..."  # Placeholder

with col2:
    st.markdown("## Risultato")
    st.write(st.session_state.get("output_text", ""))

    st.markdown("#### Lascia un feedback sul risultato")
    st.rating = st.radio("⭐️", ["⭐️", "⭐️⭐️", "⭐️⭐️⭐️", "⭐️⭐️⭐️⭐️", "⭐️⭐️⭐️⭐️⭐️"], label_visibility="collapsed", horizontal=True)
    st.text_input("Hai suggerimenti o commenti?")
    st.button("Invia feedback")
